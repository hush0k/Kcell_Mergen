import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.number_information.resolver import GraphResolver
from app.number_information.schemas import (
    NumberInformationBulkRequest,
    NumberInformationBulkResponse,
    NumberInformationRequest,
    NumberInformationResponse,
)


class OracleClientLookupService:
    def __init__(self, resolver: GraphResolver) -> None:
        self._resolver = resolver

    async def get_client_data(
            self, request: NumberInformationRequest
    ) -> NumberInformationResponse:
        return await self._resolver.resolve_fields(
            request.phone_number,
            request.fields,
            payment_date_from=request.payment_date_from,
            payment_date_to=request.payment_date_to,
        )

    async def get_clients_data(
            self, request: NumberInformationBulkRequest
    ) -> NumberInformationBulkResponse:
        return await self._resolver.resolve_fields_bulk(
            request.phone_numbers,
            request.fields,
            payment_date_from=request.payment_date_from,
            payment_date_to=request.payment_date_to,
        )

    async def export_clients_data_excel(
            self, request: NumberInformationBulkRequest
    ) -> io.BytesIO:
        response = await self.get_clients_data(request)
        columns = ["phone_number", *request.fields]
        rows = [
            {column: getattr(result, column, None) for column in columns}
            for result in response.results
        ]
        return self._to_excel(rows, columns, sheet_name="number_information")

    @staticmethod
    def _to_excel(
            rows: list[dict[str, object]], headers: list[str], sheet_name: str
    ) -> io.BytesIO:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name

        sheet.append(headers)
        for row in rows:
            sheet.append(
                [
                    value if _is_excel_native(value) else str(value)
                    for value in (row.get(header) for header in headers)
                ]
            )
        for i, header in enumerate(headers, start=1):
            max_len = max(
                len(header), *(len(str(row.get(header, ""))) for row in rows)
            ) if rows else len(header)
            sheet.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 60)

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer


def _is_excel_native(value: object) -> bool:
    return isinstance(value, (int, float, bool, date, datetime)) or value is None