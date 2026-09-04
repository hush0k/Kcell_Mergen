import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.report.repository import ReportRepository
from app.report.schemas import ReportPeriod


class ReportService:
    def __init__(self, db: AsyncSession):
        self.db = db
        self.repo = ReportRepository(db)

    async def get_list_1391(self, period: ReportPeriod) -> list[dict[str, Any]]:
        return await self.repo.get_list_1391(period.start, period.end)

    async def get_mfs_blacklist(self, period: ReportPeriod) -> list[dict[str, Any]]:
        return await self.repo.get_mfs_blacklist(period.start, period.end)

    async def get_nt_blocked(self, period: ReportPeriod) -> list[dict[str, Any]]:
        return await self.repo.get_nt_blocked(period.start, period.end)

    async def get_top10_countries(self, period: ReportPeriod) -> list[dict[str, Any]]:
        return await self.repo.get_top10_countries(period.start, period.end)

    async def export_list_1391_excel(self, period: ReportPeriod) -> io.BytesIO:
        rows = await self.get_list_1391(period)
        return self._to_excel(rows, sheet_name="list_1391")

    async def export_mfs_blacklist_excel(self, period: ReportPeriod) -> io.BytesIO:
        rows = await self.get_mfs_blacklist(period)
        return self._to_excel(rows, sheet_name="mfs_blacklist")

    async def export_nt_blocked_excel(self, period: ReportPeriod) -> io.BytesIO:
        rows = await self.get_nt_blocked(period)
        return self._to_excel(rows, sheet_name="nt_blocked")

    async def export_top10_countries_excel(self, period: ReportPeriod) -> io.BytesIO:
        rows = await self.get_top10_countries(period)
        return self._to_excel(rows, sheet_name="top10_countries")

    @staticmethod
    def _to_excel(rows: list[dict[str, Any]], sheet_name: str) -> io.BytesIO:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name

        if rows:
            headers = list(rows[0].keys())
            sheet.append(headers)
            for row in rows:
                sheet.append(
                    [
                        str(value) if not _is_excel_native(value) else value
                        for value in row.values()
                    ]
                )
            for i, header in enumerate(headers, start=1):
                max_len = max(
                    len(header), *(len(str(row.get(header, ""))) for row in rows)
                )
                sheet.column_dimensions[get_column_letter(i)].width = min(
                    max_len + 2, 60
                )
        else:
            sheet.append(["Нет данных за указанный период"])

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer


def _is_excel_native(value: object) -> bool:
    return isinstance(value, (int, float, bool, datetime)) or value is None
